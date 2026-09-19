"""
Keeps track of scanned hosts across one or more scans (an "asset
inventory") and exports that inventory to CSV or XLSX.
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from scanner import HostResult

EXPORT_COLUMNS = [
    ("IP Address", lambda h: h.ip),
    ("Hostname", lambda h: h.hostname),
    ("Status", lambda h: h.status),
    ("MAC Address", lambda h: h.mac_address),
    ("Vendor", lambda h: h.vendor),
    ("OS Guess", lambda h: h.os_name),
    ("OS Accuracy", lambda h: h.os_accuracy),
    ("Open Ports", lambda h: h.open_ports_summary),
    ("Scanned At", lambda h: h.scanned_at),
]


class AssetManager:
    """Holds the current set of known hosts, keyed by IP address."""

    def __init__(self):
        self._hosts: dict[str, HostResult] = {}

    def add_results(self, hosts: list[HostResult]) -> None:
        for host in hosts:
            self._hosts[host.ip] = host

    def clear(self) -> None:
        self._hosts.clear()

    def all_hosts(self) -> list[HostResult]:
        return sorted(self._hosts.values(), key=lambda h: _ip_sort_key(h.ip))

    def __len__(self) -> int:
        return len(self._hosts)

    # ---------- Export ----------

    def export_csv(self, path: str | Path) -> None:
        hosts = self.all_hosts()
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([label for label, _ in EXPORT_COLUMNS])
            for host in hosts:
                writer.writerow([getter(host) for _, getter in EXPORT_COLUMNS])

    def export_xlsx(self, path: str | Path) -> None:
        hosts = self.all_hosts()
        wb = Workbook()
        ws = wb.active
        ws.title = "Assets"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="C0392B", end_color="C0392B", fill_type="solid")

        for col_idx, (label, _) in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill

        for row_idx, host in enumerate(hosts, start=2):
            for col_idx, (_, getter) in enumerate(EXPORT_COLUMNS, start=1):
                ws.cell(row=row_idx, column=col_idx, value=getter(host))

        # Auto-size columns roughly, based on content width.
        for col_idx, (label, getter) in enumerate(EXPORT_COLUMNS, start=1):
            max_len = len(label)
            for host in hosts:
                val = str(getter(host))
                max_len = max(max_len, len(val))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 50)

        ws.freeze_panes = "A2"
        wb.save(path)


def _ip_sort_key(ip: str):
    try:
        return tuple(int(part) for part in ip.split("."))
    except ValueError:
        return (999, 999, 999, 999)  # non-IPv4 strings sort last
